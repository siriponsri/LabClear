import sys,json,re
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'docs/evidence/current/reader-report-v2-20261012/runtime'))
from docx import Document
from pythainlp.tokenize import word_tokenize
def th(t):return re.sub(r'[\u0e00-\u0e7f]+',lambda m:'\u200b'.join(word_tokenize(m.group(),engine='newmm',keep_whitespace=False)),t)
p=ROOT/'docs/report/LabClear_Final_Project_Report_TH.docx';d=Document(p)
for para in d.paragraphs:
 text=para.text.replace('\u200b','')
 if 'ยังไม่มีคลิปสำเร็จ' in text or 'คลิปไม่เกิน 3 นาทีเป็นข้อบังคับ' in text:
  para.text=th('คลิปสาธิตไม่เกิน 3 นาทีเป็นสิ่งส่งมอบบังคับ ขณะนี้มี LabClear_Academic_Demo_TH.webm ยาว 173.28 วินาที ตรวจเล่นจริงและตรวจเฟรมแล้ว เป็นหน้าเว็บและ API จริงในเครื่อง มีคำบรรยายไทย ไม่มีเสียง และระบุโมเดลจำลองชัดเจน จึงไม่ใช่หลักฐาน AI ใหม่ ผู้จัดทำยังต้องตรวจข้ามงานและยืนยันการส่ง Google Doc กับ Drive')
 if '1792' in text:para.text=th(text.replace('1792','1800 (ภาพเดิม 1792 และภาพผล HTTP จำลองเพิ่มเติม 8)'))
for t in d.tables:
 for row in t.rows:
  if row.cells and row.cells[0].text.startswith('R29'):row.cells[-1].text=th('มีคลิป 173.28 วินาที หน้าเว็บ/API จริง ใช้โมเดลจำลอง ไม่มีเสียง ไม่ใช่ผล AI ใหม่')
d.save(p)
p=ROOT/'docs/report/LabClear_Final_Project_Appendix_TH.docx';d=Document(p)
for para in d.paragraphs:
 if para.text.replace('\u200b','').startswith('ส่วน ช '):para.paragraph_format.page_break_before=True
actions=['บัญชีและเซสชันคนละผู้ใช้','เปิดแชตเดิมของเจ้าของเดียวกัน','เริ่มแชตใน project ที่สอง','ออกจากระบบแล้วเข้าด้วยบัญชี B','ปิด Guest แล้วเปิดเซสชันใหม่ผ่าน HTTP','เรียกแชตใหม่ในเซสชัน Guest','ล้างบริบทด้วยคำสั่งแชตใหม่ที่ระบบรองรับ','ปฏิเสธการเปิดแชตคนอื่นและส่งข้อความท้าทายสิทธิ์']
t=d.tables[-1]
for row,action in zip(t.rows[1:],actions):row.cells[1].text=th(action)
d.save(p)
p=ROOT/'docs/evidence/current/presentation-20261012/deck-payload.json';v=json.loads(p.read_text());s=v['slides']
s[7]['rows'][0][1]='เรียก AI เฉพาะจุด 9 ครั้ง';s[7]['rows'][1][1]='ผ่าน 206 ข้อในเครื่อง';s[7]['rows'][2][1]='ยอด 722/4,000 คงเดิม'
s[8]['headers'][2]='ผลที่บันทึกและตรวจ'
s[10]['numbers'][1]=['1,800','ภาพเดิม 1,792 และผล HTTP อีก 8'];s[10]['notes']+=' เพิ่มภาพผลHTTP8ภาพที่ติดป้ายชัดเจน ภาพเดิม1792byte-identical ไม่ใช่8ผลโมเดลใหม่.'
s[11]['body'][3]='คลิป 173.28 วินาทีแล้ว การตรวจข้ามงานและส่งงานยังขาด'
s[11]['notes']+=' คลิปจริง173.28วินาทีตรวจเล่นในChromiumและdecodeด้วยPlaywright FFmpeg ใช้คำบรรยายไทย ไม่มีเสียง ตัวโมเดลเป็นdouble ไม่มีprovidercallsใหม่.'
s[16]['callout']='เอกสารแนบระบุส่ง 17 ตุลาคม 2569 คลิปมีข้อจำกัดที่เปิดเผย'
v['evidence_snapshot']='722attempts/1800captures; original1792 retained'
p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf8')
print('Integrated validated clip, screenshot mode counts and compact context table.')
