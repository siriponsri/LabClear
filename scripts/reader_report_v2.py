"""Readable editable Thai report and full live-step appendix. Raw evidence unchanged."""
import sys,json,re,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'docs/evidence/current/reader-report-v2-20261012'
sys.path.insert(0,str(WORK/'runtime'))
from docx import Document
from docx.shared import Pt,Cm,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from pythainlp.tokenize import word_tokenize
WORK.mkdir(parents=True,exist_ok=True)
REPORT=ROOT/'docs/report'

def thai(text):
    return re.sub(r'[\u0e00-\u0e7f]+',lambda m:'\u200b'.join(word_tokenize(m.group(),engine='newmm',keep_whitespace=False)),str(text))

def doc(landscape=False):
    d=Document();sec=d.sections[0]
    sec.page_width=Cm(29.7 if landscape else 21);sec.page_height=Cm(21 if landscape else 29.7)
    sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.top_margin=sec.bottom_margin=Cm(1.7);sec.left_margin=sec.right_margin=Cm(1.9)
    sec.header_distance=sec.footer_distance=Cm(.7)
    for name,size in [('Normal',12),('Title',24),('Heading 1',16 if landscape else 18),('Heading 2',14),('Caption',10)]:
        s=d.styles[name];s.font.name='Tahoma';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('000000')
        s.font.bold=name.startswith('Heading')
        rp=s.element.get_or_add_rPr()
        for tag in ['szCs','bCs']:
            old=rp.find(qn('w:'+tag))
            if old is not None:rp.remove(old)
        cs=OxmlElement('w:szCs');cs.set(qn('w:val'),str(size*2));rp.append(cs)
        csbold=OxmlElement('w:bCs');csbold.set(qn('w:val'),'1' if name.startswith('Heading') else '0');rp.append(csbold)
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            s.element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
        s.element.get_or_add_rPr().append(OxmlElement('w:lang'));s.element.rPr[-1].set(qn('w:val'),'th-TH')
        s.element.rPr.rFonts.set(qn('w:eastAsia'),'Tahoma');s.element.rPr.rFonts.set(qn('w:cs'),'Tahoma')
        s.paragraph_format.line_spacing=1.3;s.paragraph_format.space_after=Pt(8)
        s.paragraph_format.widow_control=True
    for name in ['Title','Heading 1','Heading 2']:
        d.styles[name].paragraph_format.keep_with_next=True
    for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.styles['Caption'].font.color.rgb=RGBColor.from_string('55505E')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('LabClear  |  ');r.font.size=Pt(8)
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    d.core_properties.author='LabClear';d.core_properties.title='ผลการทดสอบ LabClear และประเด็นที่ต้องปรับปรุง'
    return d

def p(d,text,bold=None,style=None):
    out=d.add_paragraph(style=style);out.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=out.add_run(thai(text))
    if bold is not None:
        r.bold=bold;bc=OxmlElement('w:bCs');bc.set(qn('w:val'),'1' if bold else '0');r._r.get_or_add_rPr().append(bc)
    return out
def h(d,text,level=1):return p(d,text,style='Heading '+str(level))
def page(d,title):d.add_page_break();h(d,title)
def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,(cell,text) in enumerate(zip(t.rows[0].cells,headers)):
        cell.width=Cm(widths[i]);cell.text=thai(text)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEE8F7');cell._tc.get_or_add_tcPr().append(sh)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,(cell,text) in enumerate(zip(cells,row)):cell.width=Cm(widths[i]);cell.text=thai(str(text))
    for row in t.rows:
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            props=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
            props.append(b)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'100');x.set(qn('w:type'),'dxa');margins.append(x)
            props.append(margins)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(5);par.paragraph_format.space_before=Pt(5);par.paragraph_format.line_spacing=1.15
                for run in par.runs:
                    run.font.size=Pt(10.5);sz=OxmlElement('w:szCs');sz.set(qn('w:val'),'21');run._r.get_or_add_rPr().append(sz)
    p(d,'')
    return t

def crop(file,rect,label):
    target=WORK/(label+'.png')
    with Image.open(file) as im:im.crop(rect).save(target)
    return target
def picture(d,file,width=6.65,caption=''):
    para=d.add_paragraph();para.paragraph_format.keep_with_next=bool(caption)
    para.add_run().add_picture(str(file),width=Inches(width))
    if caption:p(d,caption,style='Caption')

base=ROOT/'docs/evidence/current/expanded-live-20261012-01'
before=crop(base/'screenshots/EXP-OCR-14-r1-view.png',(405,455,1138,558),'before-wbc-unit')
after=crop(ROOT/'docs/evidence/current/unit-literal-retest-20261012-01/screenshots/EXP-OCR-14-r1-view.png',(380,300,1165,690),'after-unit-check')
source=crop(ROOT/'docs/evidence/current/focused-unit-pixels-20261012-01/05_C_Hematology-focused-rows.png',(1215,0,1450,92),'source-unit-glyph')

def main_report():
    d=doc();p(d,'ผลการทดสอบ LabClear และสิ่งที่ต้องปรับปรุง',style='Title')
    p(d,'สำหรับผู้บริหารและอาจารย์ผู้ประเมิน  |  12 ตุลาคม 2569',style='Caption')
    h(d,'ข้อสรุปเพื่อการตัดสินใจ')
    p(d,'ยังไม่ควรเปิดใช้การแปลผลจากภาพโดยอัตโนมัติ ระบบเคยอ่านหน่วยเม็ดเลือดขาวผิด และนำหน่วยผิดไปอธิบายต่อ แม้ข้อความจะอ่านเข้าใจง่ายก็ตาม')
    p(d,'หลังปรับปรุง ระบบหยุดก่อนยืนยันข้อมูลที่ตรวจสอบหน่วยไม่ได้แล้ว การป้องกันดีขึ้น แต่ยังไม่ใช่หลักฐานว่าอ่านภาพถูกต้องครบทุกช่อง',True)
    table(d,['เรื่องที่ประเมิน','ผลที่พบ'],[
        ['การตอบคำถามและบริการ','หลายกรณีทำงานตามเกณฑ์เบื้องต้น แต่ยังพบข้อมูลที่ไม่มีหลักฐานและภาษาตอบไม่ตรงกับคำถาม'],
        ['การอ่านภาพผลตรวจ','พบหน่วยผิด ข้อมูลไม่ครบ และบางขั้นตอนจบไม่สมบูรณ์'],
        ['การป้องกันหลังแก้','การทดสอบซ้ำ 9 ครั้งหยุดข้อมูลหน่วยที่ตรวจสอบไม่ได้ ยังไม่ถือว่าอ่านภาพผ่าน'],
        ['ความพร้อมโดยรวม','ทดสอบเพิ่มเติมต่อได้เป็นกลุ่ม ยังไม่พร้อมรับรองการแปลผลจากภาพทั้งหมด']],[4.3,12.5])
    p(d,'รายงานนี้แยกผลก่อนแก้ ผลหลังแก้ และกรณีที่ยังไม่ได้ทดสอบ เพื่อให้เห็นความก้าวหน้าโดยไม่กลบข้อผิดพลาด')

    page(d,'ขอบเขตที่ทดสอบจริง')
    p(d,'ชุดหลักตามรายวิชามี 20 สถานการณ์ ประกอบด้วยข้อความ 10 ภาพ 5 และความปลอดภัย 5 สถานการณ์ ชุดเพิ่มเติมแยกจากชุดหลักอย่างชัดเจน')
    table(d,['ชุดทดสอบ','ผ่านเกณฑ์เบื้องต้น','พบปัญหา','จบไม่สมบูรณ์'],[
        ['ชุดหลัก 20 สถานการณ์','15','1','4'],
        ['ชุดเพิ่มเติมที่รันแล้ว 17 สถานการณ์','0','7','9']],[6.8,3.5,3.2,3.3])
    p(d,'อีก 1 สถานการณ์ในชุดเพิ่มเติมรอตรวจรายละเอียด และอีก 111 สถานการณ์ยังไม่ได้รันในรอบแรก การอ่านภาพเฉพาะจุดและชุดข้อความรอบต่อไปต้องรายงานแยกจากตัวเลขนี้')
    h(d,'ผลที่ไม่ต้องเรียก AI',2)
    p(d,'การตรวจความทนทานของระบบ 24 รายการผ่านทั้งหมด การทดสอบหน้าจอ 38 สถานการณ์ผ่าน 37 และมี 1 สถานการณ์ที่คลิกแล้วไม่จบภายในเวลาที่กำหนด ยังไม่ยืนยันว่าเกิดจากหน้าจอ ระบบตอบช้า หรือวิธีทดสอบ')
    p(d,'ผลเหล่านี้ช่วยยืนยันการทำงานของระบบ แต่ใช้แทนความถูกต้องของคำตอบจาก AI จริงไม่ได้')
    h(d,'ข้อมูลที่ใช้',2)
    p(d,'ภาพต้นทาง 6 ใบเป็นข้อมูลจำลอง มีรายการตรวจรวม 126 แถว ภาพที่ครอป ขยาย หรือทดสอบซ้ำยังมาจากภาพเดิม จึงไม่ใช่ตัวอย่างอิสระที่เพิ่มขึ้น และไม่ใช่การรับรองความแม่นยำทางคลินิก')

    page(d,'การอ่านหน่วยเม็ดเลือดขาวผิด')
    p(d,'พบปัญหานี้ในภาพผลเลือดสองใบ ตัวเลขที่อ่านได้เหมือนเดิม แต่เลขยกกำลังของหน่วยเปลี่ยนจาก 3 เป็น 9 เมื่อใช้ตัวหารเดิม ทำให้สเกลต่างกันหนึ่งล้านเท่า')
    p(d,'สิ่งที่ผู้ใช้เห็นก่อนแก้คือคำตอบที่ระบุว่าอยู่ในช่วงอ้างอิง ทั้งที่หน่วยในคำตอบไม่ตรงกับภาพต้นทาง ความอ่านง่ายจึงไม่รับรองความถูกต้อง')
    picture(d,source,2.2,'หน่วยจาก pixels ของภาพต้นทาง ตัวเลข 3 อยู่ตำแหน่งยกกำลัง')
    picture(d,before,6.65,'ภาพหน้าจอจริงก่อนแก้ ระบบใช้เลขยกกำลัง 9 และแสดงว่าอยู่ในช่วงอ้างอิง')
    h(d,'ผลกระทบและสิ่งที่พิสูจน์แล้ว',2)
    p(d,'หากผู้ใช้เชื่อคำตอบ อาจเข้าใจข้อมูลผลตรวจผิด ไม่มีผู้ป่วยจริงได้รับผลกระทบในชุดข้อมูลจำลองนี้')
    p(d,'ตรวจพบหน่วยผิดตั้งแต่ข้อความที่ระบบอ่านภาพส่งกลับมา จึงไม่ใช่เพียงตัวอักษรบนหน้าจอ การแก้การแสดงผลอย่างเดียวไม่เพียงพอ')

    page(d,'ข้อมูลครบและตรงกับต้นทางหรือไม่')
    h(d,'ผลตรวจบางส่วนหายจากคำอธิบาย',2)
    p(d,'ในภาพที่มีผลไทรอยด์ผิดช่วง คำตอบกล่าวถึงเฉพาะผลเลือดบางรายการที่อยู่ในช่วง จึงยังไม่ครอบคลุมข้อมูลสำคัญของภาพนั้น')
    h(d,'ชื่อภาษาไทยและอังกฤษในต้นทางไม่ตรงกัน',2)
    p(d,'ภาพน้ำตาลและปัสสาวะใช้ชื่ออังกฤษที่หมายถึงน้ำตาลหลังอดอาหาร แต่ชื่อไทยหมายถึงหลังอาหาร ระบบไม่ควรเลือกแก้ความหมายเอง ต้องแจ้งความไม่ตรงกันและให้ผู้ใช้ตรวจต้นฉบับ')
    h(d,'คำตอบบางครั้งดูครบแต่ขั้นตอนยังไม่สำเร็จ',2)
    p(d,'บางกรณีอ่านตัวเลขได้ แต่ระบบแสดงข้อความให้ลองใหม่หรือให้ตรวจภาพแทนคำอธิบายที่ครบถ้วน ต้องแยกความสำเร็จของการอ่านแต่ละช่องออกจากความสำเร็จของการตอบทั้งงาน')
    table(d,['สิ่งที่ต้องตรวจแยก','ความหมายที่ผู้อ่านใช้ประเมิน'],[
        ['ชื่อรายการและค่า','อ่านครบทุกแถว ไม่สลับค่าระหว่างแถว'],
        ['หน่วยและช่วงอ้างอิง','ตรงกับที่พิมพ์ ไม่เติมจากความรู้ทั่วไป'],
        ['ข้อความอธิบาย','กล่าวถึงข้อมูลสำคัญครบ และไม่เพิ่มข้อสรุปที่ไม่มีหลักฐาน']],[5.5,11.3])

    page(d,'การปฏิเสธผิดและการตอบคำขอที่ไม่ปลอดภัย')
    p(d,'การปฏิเสธผิดคือผู้ใช้ถามเรื่องที่ช่วยได้ แต่ระบบไม่ยอมช่วย เช่นถามว่าการตรวจวัดอะไร ส่วนการยอมทำคำขอที่ไม่ปลอดภัยคือระบบทำตามคำขอที่ควรหยุด เช่นเปิดเผยผลตรวจของคนอื่น')
    p(d,'สองปัญหานี้ต้องวัดแยกกัน ระบบที่ปฏิเสธทุกคำถามอาจหยุดคำขออันตรายได้ แต่ยังใช้งานจริงไม่ได้')
    table(d,['กลุ่มคำถามที่ตรวจด้วยคนในชุดหลัก','ผลที่พบ'],[
        ['คำขอไม่ปลอดภัย 5 สถานการณ์','หยุดหรือจำกัดคำขอได้ทั้ง 5'],
        ['คำถามทั่วไป 10 สถานการณ์','ไม่พบการปฏิเสธผิดใน 10 สถานการณ์นี้'],
        ['คำถามเพื่อความรู้สุขภาพ 2 สถานการณ์','ไม่พบการปฏิเสธผิด แต่ยังเป็นจำนวนตัวอย่างน้อย']],[9.0,7.8])
    p(d,'ผลจากชุดเล็กนี้ยังใช้สรุปกับคำถามรูปแบบใหม่ไม่ได้ ในรอบแรก ชุดความปลอดภัยเพิ่มเติมและคำสั่งแฝงยังไม่ได้รัน จึงยังไม่มีอัตราความผิดพลาดของชุดเพิ่มเติม')
    h(d,'ข้อกล่าวอ้างต้องมีหลักฐานด้วย',2)
    p(d,'คำตอบเรื่องบริการบางข้อผ่านเกณฑ์เบื้องต้น แต่พบการกล่าวถึงระยะเวลาติดต่อกลับที่แหล่งข้อมูลไม่ได้รับรอง อีกกรณีถามภาษาไทยแต่ตอบภาษาอังกฤษ จึงต้องประเมินความตรงกับแหล่งข้อมูลและความเหมาะสมของภาษาเพิ่มเติม')

    page(d,'สิ่งที่ดีขึ้นหลังแก้')
    table(d,['ก่อนแก้','หลังแก้','ข้อจำกัดที่ยังเหลือ'],[
        ['หน่วยผิดอาจถูกยืนยันและนำไปอธิบาย','หยุดก่อนยืนยันเมื่อหน่วยตรวจสอบไม่ได้','ยังไม่ได้ทำให้การอ่านภาพถูกต้องครบ'],
        ['ปุ่มยืนยันอาจใช้ค่าที่อ่านผิดเดิม','ตรวจหน่วยอีกครั้งแม้ผู้ใช้กดยืนยันทันที','ผู้ใช้ยังต้องตรวจภาพและแก้ข้อมูลที่คลาดเคลื่อน'],
        ['อาจแสดงผลว่าอยู่ในช่วงจากหน่วยผิด','ไม่ใช้หน่วยที่ยังตรวจสอบไม่ได้ตัดสินว่าอยู่ในช่วง','การตอบที่หยุดอย่างปลอดภัยไม่ใช่ผลอ่านภาพที่ผ่าน']],[5.2,5.7,5.9])
    picture(d,after,6.1,'ภาพหน้าจอจริงหลังแก้ ระบบให้ตรวจข้อมูลแทนการยืนยันหน่วยที่ยังไม่แน่ชัด')
    p(d,'การอ่านภาพเฉพาะบริเวณหน่วยกำลังทดสอบแยกต่างหาก ต้องรักษาพิกัดต้นทาง ไม่เดาเลขยกกำลัง และให้คนยืนยันก่อนใช้ข้อมูลที่แก้ไข')

    page(d,'แผนทดสอบต่อเป็นกลุ่ม')
    p(d,'ไม่ควรให้ปัญหาเฉพาะภาพหยุดการตรวจคำถามข้อความที่ไม่ขึ้นกับภาพ ชุดที่ไม่ใช้ภาพมี 62 สถานการณ์ รวม 78 ข้อความ และจะเก็บผลเป็นรอบใหม่')
    table(d,['กลุ่มที่ทดสอบต่อได้','จำนวนสถานการณ์'],[
        ['คำขอปลอดภัยและไม่ปลอดภัยที่จับคู่กัน','24'],
        ['คำถามเพื่อความรู้สุขภาพ','12'],
        ['คำสั่งแฝงและคำขอข้อมูลส่วนตัวในข้อความ','12'],
        ['คำถามที่ข้อมูลไม่ชัดเจน','6'],
        ['การแยกบริบทเมื่อเปลี่ยนแชตหรือบัญชี','8']],[12.8,4.0])
    p(d,'การอ่านหน่วยจากภาพทดสอบแบบเฉพาะจุดก่อน ไม่รันภาพเดิมซ้ำทั้งใบอย่างไม่จำกัด หากยังอ่านไม่ได้ ระบบต้องขอให้ตรวจหรือแก้ค่า แทนการเติมหน่วยเอง')
    p(d,'ภาพที่มีคำสั่งแฝง การตอบสามรูปแบบภาษา และการอธิบายผลที่ครบทุกแถวยังคงเป็นงานแยกที่ต้องทดสอบ โดยไม่ใช้ผลข้อความแทนผลภาพ')

    page(d,'งบประมาณ ข้อจำกัด และความพร้อม')
    p(d,'ตัวเลขเงินบาทที่บันทึกไว้เป็นค่าประมาณสำหรับการใช้งานที่อาจมีค่าใช้จ่ายในอนาคต ไม่ใช่ยอดใบแจ้งหนี้ เอกสารของผู้ให้บริการปัจจุบันระบุว่า Typhoon API ใช้ฟรี และ iApp ใช้ฟรีช่วง preview')
    p(d,'ก่อนทดลองอ่านหน่วยเฉพาะจุด ใช้ไป 303 ครั้งจากเพดานรวม 4,000 ครั้ง ค่าประมาณสะสม 10.84 บาท และสำรอง 0.23 บาท การหยุดรอบแรกเกิดจากหน่วยผิด ไม่ใช่เงินครบเพดาน')
    p(d,'การใช้ฟรียังมีโควตา โดยคำขอหนึ่งครั้งอาจมีหลายการตัดสินใจ ต้องนับตามรูปแบบคำขอและหยุดเมื่อถึงข้อจำกัด ไม่ใช้การเปลี่ยน key หรือรีเซ็ตประวัติเพื่อหลบโควตา')
    h(d,'ความรับผิดชอบและขอบเขตทางธุรกิจ',2)
    p(d,'พี่วัชรินรับผิดชอบ Business และศิริพลรับผิดชอบ Architecture โครงการยังไม่ได้ใช้เชิงพาณิชย์ จึงมีข้อมูลปฏิบัติงานจริงน้อย รายงานนี้ไม่ได้อ้างลูกค้า รายได้ หรือผลยืนยันทางธุรกิจที่ยังไม่มี')
    h(d,'เกณฑ์ก่อนขยายการใช้งาน',2)
    p(d,'ต้องอ่านหน่วยและข้อมูลสำคัญตรงกับต้นทาง อธิบายได้ครบ และแสดงข้อจำกัดเมื่อข้อมูลไม่ชัด ต้องไม่เปิดเผยข้อมูลของคนอื่น ไม่ทำตามคำสั่งที่ซ่อนในภาพ และไม่ปฏิเสธคำถามเพื่อความรู้โดยไม่จำเป็น')
    p(d,'ผลทั้งหมดเป็นการทดสอบซอฟต์แวร์ด้วยข้อมูลจำลอง ยังไม่ใช่การรับรองความแม่นยำในการวินิจฉัยโรคหรือความพร้อมใช้กับผู้ป่วยจริง')
    d.save(REPORT/'LabClear_Testing_Report_TH.docx')
    return d

titles={
 'Q01':'การสอบถามราคาแพ็กเกจตรวจสุขภาพ','Q02':'การสอบถามรายการตรวจในแพ็กเกจ','Q03':'การเปรียบเทียบแพ็กเกจตรวจสุขภาพ',
 'Q04':'การสอบถามสาขาที่ให้บริการ','Q05':'การยกเลิกและเปลี่ยนเวลานัด','Q06':'การขอใบเสนอราคาสำหรับองค์กร',
 'Q07':'ราคาองค์กรและเวลาติดต่อกลับ','Q08':'การอธิบายความหมายของ HbA1c','Q09':'การอธิบายความหมายของ LDL',
 'Q10':'การตอบภาษาไทยเรื่องบริการองค์กร','I01':'การอ่านผลตรวจตับ','I02':'การอ่านผลตรวจไต','I03':'การอ่านผลไขมัน',
 'I04':'การอ่านผลน้ำตาลและปัสสาวะ','I05':'การอ่านผลเม็ดเลือดและรูปร่างเซลล์',
 'S01':'การขอคำแนะนำขนาดยา','S02':'การขอผลตรวจของลูกค้าคนอื่น','S03':'การขอแก้ประวัติการเงิน',
 'S04':'การถามนโยบายคืนเงินและการร่างข้อความ','S05':'การขอเปิดเผยคำสั่งภายในระบบ'}
families=['ผลตรวจตับ','ผลตรวจไต','ผลไขมัน','ผลน้ำตาลและปัสสาวะ','ผลเลือดและรูปร่างเซลล์','ผลไทรอยด์และผลเลือด']
variants=['ภาพต้นทาง','ภาพหมุนและไม่ชัด','ภาพที่ขอบข้อมูลขาด']
for i in range(18):titles[f'EXP-OCR-{i+1:02}']='การอ่าน'+families[i//3]+'จาก'+variants[i%3]

def appendix():
    d=doc(True);p(d,'ภาพหลักฐานการทดสอบ LabClear',style='Title')
    p(d,'อ่านตามชื่อสถานการณ์และขั้นตอน ภาพทั้งหมดมาจากการทดสอบจริง คำตอบและข้อความบนหน้าจอยังคงเดิม')
    p(d,'ผลก่อนแก้และหลังแก้อยู่คนละส่วน การหยุดที่ปลอดภัยยังไม่แปลว่าอ่านภาพสำเร็จ หลักฐานเต็มทุกขั้นตอนเก็บไว้ และไม่มีการเปลี่ยนผลที่ไม่ผ่านให้เป็นผ่าน')
    map_rows=[]
    runnames=['expanded-live-20261012-01','unit-fidelity-retest-20261012-02','unit-fidelity-retest-20261012-03','unit-fidelity-retest-20261012-04','unit-fidelity-source-probe-20261012-01','unit-literal-retest-20261012-01']
    scores={s['id']:s for s in json.loads((base/'scored.json').read_text())}
    status={'PASS':'ผ่านเกณฑ์เบื้องต้น','FAIL':'พบข้อผิดพลาด','ERROR':'ขั้นตอนจบไม่สมบูรณ์','REVIEW_PENDING':'รอตรวจรายละเอียด'}
    stepnames={'view':'ผลที่ผู้ใช้เห็น','full':'หน้าจอและบทสนทนาทั้งหมด','review':'รายละเอียดที่อ่านจากภาพ'}
    for phase,run in enumerate(runnames):
        folder=ROOT/'docs/evidence/current'/run
        cases=[json.loads(s) for s in (folder/'cases.jsonl').read_text().splitlines() if s.strip()]
        for c in cases:
            title=titles.get(c['id'],'การตรวจคำตอบตามสถานการณ์')
            period='ก่อนแก้การตรวจหน่วย' if phase==0 else 'หลังแก้การตรวจหน่วย รอบที่ '+str(phase)
            verdict=status.get(scores.get(c['id'],{}).get('verdict'),'รอตรวจรายละเอียด') if phase==0 else 'หยุดให้ตรวจหน่วย ยังไม่ถือว่าอ่านภาพผ่าน'
            if c['id'] in {'Q01','Q02','Q07','Q10'} and phase==0:verdict+=' และมีประเด็นที่ต้องตรวจเพิ่มเติม'
            for shot in c.get('screenshots',[]):
                if not shot.get('file'):continue
                file=folder/shot['file'];image=Image.open(file)
                # Crop only empty left navigation on ordinary result views; all context stays in full view.
                left=280 if image.width==1280 else 0
                image=image.crop((left,0,image.width,image.height))
                width=image.width;panel_h=int(width*.52)
                for start in range(0,image.height,panel_h):
                    end=min(start+panel_h,image.height)
                    if end-start<65 and start:continue
                    panel=WORK/f"panel-{len(map_rows)+1:04}.png";image.crop((0,start,width,end)).save(panel)
                    d.add_page_break();h(d,title)
                    p(d,period+'  |  '+stepnames.get(shot.get('step'),'ภาพขั้นตอนที่ทดสอบ')+'  |  '+verdict,style='Caption')
                    picture(d,panel,9.7,'ภาพหน้าจอจริง ส่วนที่ '+str(start//panel_h+1)+' เนื้อหาภาพไม่ได้แก้ไข ขยายเฉพาะการจัดหน้าเพื่อให้อ่านได้')
                    map_rows.append({'reader_title':title,'phase':period,'case_id':c['id'],'run_id':run,'original_file':str(file),'original_sha256':shot.get('sha256'),'step':shot.get('step'),'panel_crop_xyxy':[left,start,left+width,end],'panel_file':str(panel),'result':verdict})
    d.save(REPORT/'LabClear_Testing_Appendix_TH.docx')
    (WORK/'reader-evidence-mapping.json').write_text(json.dumps(map_rows,ensure_ascii=False,indent=2),encoding='utf-8')
    return len(map_rows)

main_report()
if '--main-only' not in sys.argv:print('appendix panels',appendix())
print('Reader report DOCX authored with editable headings paragraphs tables and Thai word breaks')
