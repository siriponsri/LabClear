import sys,re,json
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'docs/evidence/current/reader-report-v2-20261012/runtime'))
from docx import Document
from docx.shared import RGBColor,Pt
from docx.oxml.ns import qn
from pythainlp.tokenize import word_tokenize
def th(t):return re.sub(r'[\u0e00-\u0e7f]+',lambda m:'\u200b'.join(word_tokenize(m.group(),engine='newmm',keep_whitespace=False)),t)
p=ROOT/'docs/report/LabClear_Final_Project_Report_TH.docx';d=Document(p)
replace={'session':'บัญชีและช่วงการใช้งาน','CSRF':'ป้องกันคำสั่งปลอมผ่านบัญชีผู้ใช้','citation_invalid':'ข้อผิดพลาดแหล่งอ้างอิงไม่ตรง','isolation_action':'การสลับบัญชีหรือแชตเพื่อทดสอบแยกข้อมูล','gate receipts':'รายการรับเข้าที่ใช้ตรวจนับการเรียก AI','shadow future estimate':'ประมาณการค่าใช้จ่ายในอนาคตแยกจากยอดจริง','shadow estimates':'ประมาณการค่าใช้จ่ายแยกจากยอดจริง','demo fixture':'ข้อมูลจำลองสำหรับทดสอบ','lexical':'ค้นจากคำที่ตรงกัน','vector database':'ฐานค้นจากความใกล้เคียงเชิงความหมาย','input safety':'การตรวจคำขอที่รับเข้า','input guard':'ด่านตรวจคำขอ','provider attempts':'จำนวนครั้งที่เรียกผู้ให้บริการ','whole-case automatic retry':'การรันทั้งกรณีซ้ำอัตโนมัติ','workflow ceiling':'เพดานต่อขั้นตอนงาน','Monetary ledger settled estimate':'ยอดค่าใช้จ่ายที่ลงบัญชีแบบประมาณ','reserved':'ยอดที่กันไว้','admissions':'รายการที่รับเข้าแล้ว','reservation':'ยอดที่กันไว้','interruption':'การขาดตอน','recovery':'ข้อความช่วยให้ผู้ใช้ดำเนินการต่อ','numeric range':'ช่วงตัวเลข','Printed-flag classifier':'การอ่าน flag ที่พิมพ์ในภาพ','Final-answer quality':'คุณภาพคำตอบสุดท้าย','source completeness':'ความครบถ้วนของแหล่งต้นทาง','false negatives':'คำขออันตรายที่หลุดการตรวจ','false positives':'คำขอปลอดภัยที่ถูกหยุดผิด','cross-review':'การตรวจข้ามงานของผู้จัดทำ','mock':'โมเดลจำลอง','offline':'ทดสอบในเครื่องโดยไม่เรียก AI','live':'ทดสอบกับผู้ให้บริการจริง'}
for para in d.paragraphs:
 plain=para.text.replace('\u200b','')
 if not plain or plain.startswith('หลักฐานหลัก:'):continue
 # Keep code paths and source names in the evidence register, not in the prose.
 for a,b in replace.items():plain=re.sub(r'(?<![\w/])'+re.escape(a)+r'(?![\w/])',lambda m:b,plain,flags=re.I)
 plain=plain.replace('prior conservative 24 + new','ยอดเดิมที่นับเผื่อไว้ 24 ครั้ง รวมกับยอดใหม่')
 if 'ต้นแบบมี FastAPI' in plain:
  plain='หน้าเว็บส่งข้อความให้ FastAPI ซึ่งตรวจบัญชีและสิทธิ์เจ้าของข้อมูลก่อนทำงาน ระบบค้นจากคำที่ตรงกับเอกสารด้วยวิธี BM25 แล้วเลือกเครื่องมือและแหล่งที่อนุมัติ ผู้ให้บริการ Typhoon ช่วยอ่านภาพและเขียนข้อความ ส่วน iApp ตรวจความปลอดภัย บัญชีและสถานะงานเก็บใน SQLite การมีขั้นตรวจคำตอบไม่ได้รับรองว่าคำตอบถูกเสมอ'
 if 'รอบข้อความ62' in plain.replace(' ','') and 'ผล' in plain:pass
 para.text=th(plain)
# The front summary gets the current interpretation next to the historical automation.
for t in d.tables:
 for row in t.rows:
  for cell in row.cells:
   plain=cell.text.replace('\u200b','')
   if '62 กรณี' in plain and 'ข้อความ' in plain:pass
   if '39' in plain and 'REVIEW_PENDING' in plain:
    cell.text=th('ผลอัตโนมัติเดิม: รอตรวจ 39 ไม่ผ่าน 11 ผิดพลาด 12\nตรวจคำตอบแล้ว: ผ่าน 9 บางส่วน 20 ไม่ผ่าน 22 ผิดพลาด 11')
   elif 'Session / CSRF' in plain or 'session / CSRF' in plain:cell.text=th('ตรวจบัญชี ป้องกันคำสั่งปลอม และตรวจสิทธิ์เจ้าของข้อมูล')
   elif 'gate receipts' in plain:cell.text=th(plain.replace('gate receipts','รายการรับเข้าที่ตรวจนับ AI'))
for para in d.paragraphs:
 if 'คลิปไม่เกิน 3 นาทีเป็นข้อบังคับ' in para.text.replace('\u200b',''):
  para.text=th('คลิปไม่เกิน 3 นาทีเป็นข้อบังคับ มีคลิปหน้าเว็บและ API จริงในเครื่องที่กำลังตรวจเล่นและวัดเวลาของไฟล์ ใช้คำบรรยายไทยและโมเดลจำลองติดป้ายชัดเจน ไม่แทนการทดสอบ AI จริง การตรวจข้ามงาน การส่ง Google Doc และ Drive ยังรอผู้จัดทำ')
d.save(p)
p=ROOT/'docs/presentation/LabClear_Academic_Demo_Storyboard_TH.docx';d=Document(p)
for st in d.styles:
 if st.name in ['Title','Heading 1','Heading 2','Normal']:
  st.font.color.rgb=RGBColor(0,0,0);st.font.bold=st.name!='Normal'
  rp=st.element.get_or_add_rPr();color=rp.find(qn('w:color'))
  if color is not None:
   for key in list(color.attrib):
    if 'theme' in key.lower():del color.attrib[key]
for e in list(d.styles.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
for para in d.paragraphs:
 for r in para.runs:
  if para.style.name in ['Title','Heading 1','Heading 2']:r.font.color.rgb=RGBColor(0,0,0)
d.save(p)
print('Main narrative simplified and storyboard headings made black; no raw answers edited.')
