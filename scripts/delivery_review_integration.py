"""Integrate separate semantic outcomes and offline boundary evidence into reports."""
import sys,json,re
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'docs/evidence/current/reader-report-v2-20261012/runtime'))
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pythainlp.tokenize import word_tokenize
EV=ROOT/'docs/evidence/current/full-course-report-20261012'
reviews=json.loads((EV/'text-outcome-review.json').read_text());summary=json.loads((EV/'semantic-review-summary.json').read_text())
def th(t):return re.sub(r'[\u0e00-\u0e7f]+',lambda m:'\u200b'.join(word_tokenize(m.group(),engine='newmm',keep_whitespace=False)),t)
def para(d,t,style=None):return d.add_paragraph(th(t),style)
def table(d,headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 for c,text,w in zip(t.rows[0].cells,headers,widths):c.text=th(text);c.width=Cm(w)
 h=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(h)
 for data in rows:
  for c,text,w in zip(t.add_row().cells,data,widths):c.text=th(str(text));c.width=Cm(w)
 for row in t.rows:
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.keep_together=True;p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.name='TH SarabunPSK';r.font.size=Pt(14);rp=r._r.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    for r in p.runs:
     rf=r._r.get_or_add_rPr().find(qn('w:rFonts'))
     for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'TH SarabunPSK')
 return t
p=ROOT/'docs/report/LabClear_Final_Project_Report_TH.docx';d=Document(p)
for para0 in d.paragraphs:
 text=para0.text.replace('\u200b','')
 if '39 REVIEW_PENDING' in text or ('39' in text and 'pending' in text.lower()):
  # Preserve historical automated figures, append current distinction.
  para0.add_run(th(' การตรวจคำตอบที่บันทึกไว้เพิ่มเติมให้ผล PASS 9, PARTIAL 20, FAIL 22 และ ERROR 11 เป็นคนละเกณฑ์กับผลอัตโนมัติเดิม'))
 if 'audit พบ context 8' in text:
  para0.text=th('context เดิม 8 กรณีไม่ได้สลับบัญชีหรือแชตจริง จึงไม่ใช่หลักฐานแยกข้อมูล การทดสอบ HTTP เพิ่มเติมบนฐานข้อมูลจำลองใหม่ใช้การสลับบัญชี แชต project ออกจากระบบ เข้าใหม่ และปิดเซสชัน Guest จริง ผ่าน 8 ข้อโดยตรวจบริบทที่ส่งให้ agent double ไม่เรียก AI และไม่รับรองว่าระบบกันข้อมูลรั่วได้ทุกสถานการณ์')
# Remove only our short appended readiness paragraph; existing clinical limitation stays.
for para0 in list(d.paragraphs):
 if para0.text.startswith('คลิปสาธิตต้นแบบจริงไม่เกิน 3 นาทีเป็นสิ่งส่งมอบ'):para0._p.getparent().remove(para0._p)
para(d,'ผลตรวจคำตอบเพิ่มเติมและขอบเขตที่ยืนยันได้','Heading 2')
para(d,'อ่านคำตอบที่บันทึกไว้ 62 กรณีโดยเทียบแหล่งข้อมูลธุรกิจและความรู้ที่ระบบใช้ เก็บผลอัตโนมัติเดิมแยกไว้ ไม่มีการเรียกผู้ให้บริการใหม่ ผล PASS จำกัดเฉพาะเจตนา นโยบาย และข้ออ้างที่ตรวจ ส่วน PARTIAL ยังมีข้อจำกัดตามเหตุผลในภาคผนวก ไม่ใช่ผลทางคลินิกที่ได้รับการรับรอง')
table(d,['ผลตรวจคำตอบ','จำนวน','ความหมาย'],[['PASS',9,'ตอบเจตนาและขอบเขตที่ตรวจได้'],['PARTIAL',20,'ตอบบางส่วนหรือหยุดปลอดภัย แต่ยังมีข้อจำกัด'],['FAIL',22,'พบข้ออ้างไม่รองรับ นิยามผิด หรือปฏิเสธคำถามดี'],['ERROR',11,'ไม่มีผลครบเวิร์กโฟลว์']], [3.1,1.9,12.4])
para(d,'ใน 39 กรณีที่ระบบให้ REVIEW_PENDING เดิม มี PASS 9, PARTIAL 15 และ FAIL 15 จากการตรวจรอบนี้ ความไม่แน่ชัดของ PARTIAL ไม่ถูกเปลี่ยนเป็น PASS ส่วนข้อความที่สามของ CONTEXT05 ไม่เกิดขึ้นเพราะข้อความที่สองผิดพลาดตามเงื่อนไขหยุดของ harness จึงไม่สร้างผลย้อนหลัง')
para(d,'คลิปไม่เกิน 3 นาทีเป็นข้อบังคับ ขณะนี้อยู่ระหว่างบันทึกและตรวจเล่นจริง หากผ่านจะเป็นคลิปหน้าเว็บและ API จริงในเครื่อง มีคำบรรยายไทยและโมเดลจำลองที่ติดป้ายชัดเจน ไม่แทนการทดสอบ AI จริง การตรวจข้ามงาน การส่ง Google Doc และ Drive ยังรอผู้จัดทำ')
d.save(p)
p=ROOT/'docs/report/LabClear_Final_Project_Appendix_TH.docx';d=Document(p);d.add_page_break();para(d,'ส่วน ฉ ผลตรวจเชิงความหมายของคำตอบเดิม','Heading 1')
para(d,'ตรวจจากผลที่บันทึกแล้ว ไม่รัน provider ซ้ำ และไม่แก้ผลอัตโนมัติหรือคำตอบเดิม ตารางนี้แยกผลคำตอบออกจากการจำแนกคำขออันตราย ข้อผิดพลาดการทำงานไม่ถูกนับเป็น TN หรือการปฏิเสธสำเร็จ PARTIAL ระบุช่องว่างที่ยังไม่รับรองทางคลินิก')
table(d,['Case ID','อัตโนมัติเดิม','คำตอบ','เหตุผลจากการตรวจ'],[[r['id'],r['automated_verdict'],r['semantic_outcome'],r['rationale_th']] for r in reviews],[3.5,2.7,1.8,9.4])
para(d,'แหล่งข้อมูลที่ตรวจ: business_data/policies.json, catalog.json, branches.json และ knowledge/evidence/catalog.json ผลระดับกรณีพร้อม hash ของคำตอบอยู่ใน text-outcome-review.json; raw cases.jsonl คงเดิม')
para(d,'ส่วน ช การทดสอบแยกบริบทผ่าน HTTP จริง','Heading 1')
ctx=json.loads((ROOT/'docs/evidence/current/context-isolation-offline-20261012-01/results.json').read_text())
table(d,['Case ID','การกระทำจริง','ผล'],[[r['id'],r['action'],r['actual_outcome']] for r in ctx['cases']],[3.5,11.4,2.5])
para(d,'ใช้ source ZIP ที่แก้แล้วและฐานข้อมูลจำลองใหม่ ตรวจสิทธิ์และบริบทที่ส่งให้ agent double ซึ่งบันทึกเฉพาะ canary ที่สร้างเพื่อทดสอบ การปิด/เปิด Guest ทำผ่าน HTTP เทียบลำดับหน้าเว็บ ไม่ใช่หลักฐาน browser reload การทดสอบนี้ทำ provider calls 0 และคงยอดรวมเดิม 722; ไม่วัดการต้านคำสั่งของโมเดลจริง')
d.save(p)
payload=ROOT/'docs/evidence/current/presentation-20261012/deck-payload.json';v=json.loads(payload.read_text());s=v['slides']
s[7]['rows'][0][0]='ตรวจหน่วยก่อนยืนยัน';s[7]['rows'][1][0]='จับคู่เซลล์กับแถว';s[7]['rows'][2][0]='นับทุกครั้งที่เรียก AI'
s[8]['rows'][1][2]='ตรวจคำตอบ: ผ่าน 9 บางส่วน 20 ไม่ผ่าน 22 ผิดพลาด 11';s[8]['callout']='ผลอัตโนมัติเดิมยังอยู่ การตรวจคำตอบแยกจากการจำแนกความปลอดภัย'
s[11]['body'][2]='ตรวจขอบเขต HTTP ผ่าน 8 ข้อ ยังไม่วัดโมเดลจริง';s[11]['body'][3]='คลิปกำลังตรวจเล่นจริง การตรวจข้ามงานและส่งงานยังขาด'
s[15]['rows'][1][0]='คำสั่งข้อความ';s[15]['headers'][1]='ข้อมูลอันตราย / ควบคุม'
s[8]['notes']+=' การตรวจคำตอบเดิมเพิ่มเติม62กรณี: PASS9 PARTIAL20 FAIL22 ERROR11. ใน39REVIEW_PENDINGเดิมเป็นPASS9 PARTIAL15 FAIL15 คงผลอัตโนมัติและข้อจำกัดไว้ ไม่ใช่ clinical validation.'
s[11]['notes']+=' ขอบเขตHTTPจริงเพิ่มเติม8PASSด้วยagent doubleและฐานข้อมูลใหม่ ไม่มีprovidercalls. Guestclose/reloadเป็นHTTPsequenceไม่ใช่browserreload.'
payload.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'semantic_cases_added':62,'offline_context_cases':8,'provider_calls':0}))
