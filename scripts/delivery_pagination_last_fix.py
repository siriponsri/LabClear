import sys,re
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'docs/evidence/current/reader-report-v2-20261012/runtime'))
from docx import Document
from pythainlp.tokenize import word_tokenize
def th(t):return re.sub(r'[\u0e00-\u0e7f]+',lambda m:'\u200b'.join(word_tokenize(m.group(),engine='newmm',keep_whitespace=False)),t)
p=ROOT/'docs/report/LabClear_Final_Project_Report_TH.docx';d=Document(p)
for i,para in enumerate(d.paragraphs):
 if para.text.replace('\u200b','').startswith('4. ระบบ'):
  d.paragraphs[i+1].text=th('หน้าเว็บส่งคำถามให้ FastAPI ตรวจบัญชีและสิทธิ์ แล้วค้นเอกสารจากคำที่ตรงกัน (BM25) ก่อนเรียก AI เพื่อเขียนและตรวจคำตอบ สถานะงานเก็บใน SQLite ผลทดสอบยังพบข้อผิดพลาดที่เปิดเผยในรายงาน')
for table in d.tables[:2]:
 for row in table.rows:
  key=row.cells[0].text.replace('\u200b','').replace(' ','')
  if 'ข้อความ' in key and '62' in key:
   row.cells[0].text=th('ตรวจคำตอบเพิ่มเติม 62 กรณี');row.cells[1].text=th('ผ่าน 9 บางส่วน 20 ไม่ผ่าน 22 ผิดพลาด 11')
d.save(p);print('Shortened architecture introduction to avoid an isolated continuation page.')
