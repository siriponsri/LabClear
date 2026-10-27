"""Four targeted artifact repairs only. No app/provider imports or test execution."""
import sys,json,re,zipfile,hashlib
from pathlib import Path
from datetime import datetime,timezone
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'docs/evidence/current/reader-report-v2-20261012/runtime'))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pythainlp.tokenize import word_tokenize
EV=ROOT/'docs/evidence/current/focused-remnants-20261012';EV.mkdir(exist_ok=True)
def sha(b):return hashlib.sha256(b).hexdigest()
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def save(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf8')
def th(t):return re.sub(r'[\u0e00-\u0e7f]+',lambda m:'\u200b'.join(word_tokenize(m.group(),engine='newmm',keep_whitespace=False)),t)
protected=['docs/report/LabClear_Source_Code_20261012.zip','docs/presentation/LabClear_Academic_Demo_TH.webm','docs/presentation/LabClear_Academic_Demo_Timings.json','docs/evidence/current/independent-text-20261012-01/cases.jsonl']
if not (EV/'protected-before.json').exists():save(EV/'protected-before.json',{n:sha((ROOT/n).read_bytes()) for n in protected})
p=ROOT/'docs/report/LabClear_Final_Project_Report_TH.docx'
if not (EV/(p.name+'.before')).exists():(EV/(p.name+'.before')).write_bytes(p.read_bytes())
d=Document(p);row_found=0
for t in d.tables:
 for row in t.rows:
  if row.cells[0].text.replace('\u200b','')=='ความพร้อมส่งงาน':
   row.cells[1].text=th('มีคลิป 173.28 วินาทีแล้ว ใช้โมเดลจำลองและคำบรรยายไทย ไม่มีเสียง รอเจ้าของตรวจความเหมาะสมก่อนส่ง การตรวจข้ามงานและส่งบน Drive ยังไม่เสร็จ');row_found+=1
assert row_found==1
replaced=[]
for para in d.paragraphs:
 t=para.text.replace('\u200b','')
 if 'exact field accuracy' in t:
  para.text=th('ตรวจการอ่านภาพทีละช่องว่าชื่อรายการ ค่า หน่วย และช่วงอ้างอิงตรงกับต้นทางหรือไม่ พร้อมตรวจว่ามีรายการครบทั้งใบ ยอมรับการแปลงหน่วยเฉพาะเมื่อแทนปริมาณเดียวกัน ไม่ให้คะแนนจากการเดา การเทียบช่วงใช้ตัวเลขที่พิมพ์ในรายงาน ชุดนี้มี 40 แถวนอกช่วงและ 71 แถวในช่วง แต่ยังไม่มีผลอ่านทุกแถวที่ตรวจยืนยัน จึงยังคำนวณจำนวนทำนายถูกและผิดไม่ได้');replaced.append('exact_field_accuracy')
 elif 'claim-level gold review' in t:
  para.text=th('การอ่านเครื่องหมาย H/L/HH ที่พิมพ์ต้องแยกจากการเทียบช่วงตัวเลข มี 42 แถวที่พิมพ์เครื่องหมายและ 84 แถวที่ไม่มี ส่วนการตรวจข้อกล่าวอ้างในคำตอบต้องให้ผู้ประเมินเทียบทุกข้อกับแหล่งข้อมูลก่อน แล้วระบุว่าข้อใดมีหลักฐานหรือไม่มีหลักฐาน ปัจจุบันพบตัวอย่างที่ไม่รองรับ แต่ยังไม่ตรวจครบทุกข้อ จึงไม่รายงานอัตราความถูกต้องทั้งคำตอบจากการมีแหล่งอ้างอิงเพียงอย่างเดียว');replaced.append('claim_level_gold_review')
 elif 'gate เพดาน' in t:
  para.text=th('การทดสอบด้วยโมเดลจำลองตรวจเส้นทาง HTTP การเก็บข้อมูล สิทธิ์บัญชี การป้องกันคำสั่งปลอม และการเปลี่ยนสถานะงาน โดยไม่วัดคุณภาพ AI เมื่อส่งซ้ำ หมดเวลา ยกเลิก หรือได้ผลที่อ่านไม่ได้ ต้องนับทุกครั้งที่เรียกผู้ให้บริการ รอบนี้ไม่รันทั้งกรณีซ้ำอัตโนมัติ และกำหนดเพดาน 12 ครั้งต่อขั้นตอนงานเพื่อกันการเรียกวน ข้อผิดพลาดไม่นับเป็นการจำแนกคำขอถูกหรือเป็นหลักฐานว่าความเป็นส่วนตัวผ่าน');replaced.append('gate_narrative')
for t in d.tables:
 for row in t.rows:
  for cell in row.cells:
   text=cell.text.replace('\u200b','')
   if 'physical scale' in text:
    cell.text=th('ตรวจว่าตัวเลขและหน่วยยังแทนปริมาณเดิม เช่น เลขยกกำลังต้องไม่เปลี่ยน หากยืนยันหน่วยไม่ได้ให้หยุด การรันทดสอบเฉพาะจุดกับ AI 9 ครั้งหยุดข้อมูลหน่วยที่ตรวจไม่ได้ ยังไม่ใช่ผลอ่านภาพผ่าน');replaced.append('physical_scale')
   elif 'ambiguous spans' in text:
    cell.text=th('รองรับตาราง 5–7 คอลัมน์และเก็บชื่อรายการ ค่า หน่วยแยกกัน หากตำแหน่งเซลล์ทำให้จับคู่ข้อมูลไม่ได้ให้หยุด การทดสอบในเครื่องผ่าน 206 ข้อ ไม่แทนความถูกต้องของ AI ทุกภาพ');replaced.append('ambiguous_spans')
   elif 'atomic admission gate' in text:
    cell.text=th('ลงบัญชีก่อนส่งคำขอเพื่อไม่ให้คำขอพร้อมกันเกินเพดาน นับครั้งที่ล้มเหลวและยกเลิกด้วย ตรวจระบบนับ 9 ข้อและโควตาคำถาม 8 ข้อ ยอดในรอบนั้นคือ 635/4000 ประมาณค่าใช้จ่ายเฉพาะรอบที่ยืนยันว่าฟรี โดยไม่ล้างประวัติบัญชีจริง');replaced.append('gate_table')
d.save(p)
# Main Markdown companion only; preserve all other deliverables.
parts=[]
from docx.table import Table
from docx.text.paragraph import Paragraph
for el in d.element.body:
 if el.tag==qn('w:p'):
  para=Paragraph(el,d);text=para.text.replace('\u200b','').replace('\u2060','')
  if text:parts.append(('# ' if para.style.name=='Title' else '## ' if para.style.name=='Heading 1' else '### ' if para.style.name=='Heading 2' else '')+text)
 elif el.tag==qn('w:tbl'):
  rows=[[c.text.replace('\u200b','').replace('\u2060','').replace('|','\\|').replace('\n','<br>') for c in row.cells] for row in Table(el,d).rows]
  parts.append('\n'.join(['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']+['| '+' | '.join(row)+' |' for row in rows[1:]]))
(ROOT/'docs/report/LabClear_Final_Project_Report_TH.md').write_text('\n\n'.join(parts)+'\n',encoding='utf8')

# Add an explicit ZIP member to every current screenshot record, retaining source paths.
archive=ROOT/'docs/report/LabClear_All_Screenshots_20261012.zip'
if not (EV/'screenshots-archive-before.json').exists():save(EV/'screenshots-archive-before.json',{'sha256':sha(archive.read_bytes())})
with zipfile.ZipFile(archive) as z:content={n:z.read(n) for n in z.namelist()}
original_pngs={n:sha(b) for n,b in content.items() if n.lower().endswith('.png')}
def bind(record):
 source=str(record.get('file') or record.get('source','')).replace('\\','/')
 preferred=[source,str(record.get('archive_member',''))]
 if record.get('mode')=='OFFLINE_HTTP_RESULT_VIEW':preferred.insert(0,'context-isolation-offline-20261012-01/'+Path(source).name)
 chosen=next((n for n in preferred if n in content and n.lower().endswith('.png')),None)
 assert chosen is not None,source
 assert sha(content[chosen])==record['sha256'],chosen
 record['archive_member']=chosen
 return record
masterpath=ROOT/'docs/evidence/current/full-course-report-20261012/all-screenshots-manifest.json';master=load(masterpath)
master['captures']=[bind(r) for r in master['captures']];assert len(master['captures'])==1800;save(masterpath,master)
contextpath=ROOT/'docs/evidence/current/context-isolation-offline-20261012-01/screenshots-manifest.json';context=[bind(r) for r in load(contextpath)];save(contextpath,context)
content['ALL_SCREENSHOTS_CURRENT_MANIFEST.json']=json.dumps(master,ensure_ascii=False,indent=2).encode()
content['OFFLINE_CONTEXT_SCREENSHOTS_MANIFEST.json']=json.dumps(context,ensure_ascii=False,indent=2).encode()
for name in ['SCREENSHOTS_MANIFEST.json']:
 if name in content:
  data=json.loads(content[name]);records=data.get('captures',[]) if isinstance(data,dict) else data
  for record in records:bind(record)
  content[name]=json.dumps(data,ensure_ascii=False,indent=2).encode()
temp=archive.with_suffix('.focused.zip')
with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in content.items():z.writestr(n,b)
temp.replace(archive)
with zipfile.ZipFile(archive) as z:
 failures=[]
 for r in master['captures']:
  if r['archive_member'] not in z.namelist() or sha(z.read(r['archive_member']))!=r['sha256']:failures.append(r['file'])
 assert not failures
 assert all(sha(z.read(n))==h for n,h in original_pngs.items())
save(EV/'focused-remnants-checks.json',{'at':datetime.now(timezone.utc).isoformat(),'executive_row_replaced':row_found,'jargon_replacements':replaced,'archive_manifest_entries':1800,'direct_member_lookup_failures':failures,'all_1800_png_bytes_unchanged':True,'source_video_raw_protected':{n:sha((ROOT/n).read_bytes())==h for n,h in load(EV/'protected-before.json').items()},'new_provider_calls':0,'new_test_cases':0,'slide8_native_render':'PENDING','report_render':'PENDING'})
print(json.dumps({'executive_row':row_found,'jargon_replacements':replaced,'direct_manifest_entries_verified':1800,'PNG_bytes_unchanged':True,'provider_calls':0}))
